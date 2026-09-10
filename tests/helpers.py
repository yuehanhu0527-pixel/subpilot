"""测试辅助工具：可预测的假 embedder 与文档 fixture 生成器。

只有测试会 import 这个文件——生产代码完全不依赖它。
"""

from pathlib import Path


class FakeEmbedder:
    """把已知单词映射到固定正交向量；完全未知的文本落到固定的"无关"向量。

    用于快速、确定性地验证检索与阈值行为（不下载/运行真实模型）。
    """

    def __init__(self, vocab: dict[str, list[float]], dim: int):
        self._vocab = vocab
        self._dim = dim
        # 没有任何已知词的文本 → 与词表正交的向量（避免零向量引发除零）
        self._unknown = [0.0] * dim
        self._unknown[-1] = 1.0

    def encode(self, texts: list[str]) -> list[list[float]]:
        out = []
        for text in texts:
            vec = [0.0] * self._dim
            hits = 0
            for word in text.lower().split():
                word = word.strip(".,!?;:'\"()")
                if word in self._vocab:
                    vec = [a + b for a, b in zip(vec, self._vocab[word])]
                    hits += 1
            if hits == 0:
                vec = list(self._unknown)
            else:
                # 多词平均但不归一化 → 与单主题文档的相似度 < 1，可测阈值行为
                vec = [x / hits for x in vec]
            out.append(vec)
        return out


def make_text_pdf(path: Path, lines: list[str]) -> Path:
    """手写一个最小合法的单页 PDF，每行一条文本（Helvetica 字体）。"""
    def esc(s: str) -> str:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    cmds = []
    y = 720
    for line in lines:
        cmds.append(f"BT /F1 12 Tf 72 {y} Td ({esc(line)}) Tj ET")
        y -= 18
    stream = "\n".join(cmds).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n").encode()
    path.write_bytes(bytes(out))
    return path


def make_docx(path: Path, heading: str | None, body: str) -> Path:
    """生成一个含可选标题 + 正文段落的 DOCX。"""
    import docx

    doc = docx.Document()
    if heading:
        doc.add_heading(heading, level=1)
    doc.add_paragraph(body)
    doc.save(str(path))
    return path
